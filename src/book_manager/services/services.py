"""Servicios con la lógica de negocio de Book Manager."""
import abc
import csv
import datetime
import functools
import inspect
import logging
import os
import re
import threading
import unicodedata
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import (Any, Callable, Dict, Generic, List, NamedTuple, Optional, Sequence, Tuple,
                    TypeVar)

import requests

from book_manager.database.connection import RAIZ_PROYECTO, cargar_variables_entorno
from book_manager.entities.entities import (
    Auditoria, CotizacionDolar, Editorial, EntidadBase, Genero, Libro, Moneda,
    Precio, PrecioWeb, Stock, TipoCotizacion)
from book_manager.repositories.repositories import (
    RepositorioAuditoria, RepositorioCotizacionDolar, RepositorioDB, RepositorioEditorial,
    RepositorioGenero, RepositorioLibro, RepositorioMoneda, RepositorioPrecio,
    RepositorioPrecioWeb, RepositorioStock, RepositorioTipoCotizacion)

T = TypeVar('T', bound=EntidadBase)

CODIGO_ARS = "ARS"
CODIGO_USD = "USD"
CARPETA_EXPORTACION = os.path.join(RAIZ_PROYECTO, "exportaciones")


# ---------------------------------------------------------------- auditoría
_contexto = threading.local()  # cuántas operaciones auditadas están en curso (anidadas)
_servicio_auditoria: Optional["ServicioAuditoria"] = None


def configurar_auditoria(servicio: Optional["ServicioAuditoria"]) -> None:
    """Indica qué servicio usa el decorador @auditar para registrar las operaciones.

    Con None se desactiva la auditoría (por ejemplo, al cerrar la conexión).
    """
    global _servicio_auditoria
    _servicio_auditoria = servicio


class _SinFaltantes(dict):
    """Diccionario para formatear nombres de acción: lo que no existe se reemplaza por ''."""

    def __missing__(self, clave: str) -> str:
        """Devuelve una cadena vacía para las claves que no están."""
        return ""


def _resumir(valor: Any) -> str:
    """Describe un valor en pocas palabras para guardarlo en los detalles de la auditoría."""
    if isinstance(valor, (list, tuple, set, frozenset, dict)) and not hasattr(valor, "_fields"):
        return f"{len(valor)} elementos"
    texto = str(valor)
    if isinstance(valor, EntidadBase):
        texto = f"{texto} (id {valor.id})"
    return texto if len(texto) <= 200 else texto[:197] + "..."


def _describir_argumentos(argumentos: Dict[str, Any]) -> str:
    """Arma el texto 'nombre=valor' con los argumentos de una operación."""
    return ", ".join(f"{nombre}={_resumir(valor)}" for nombre, valor in argumentos.items()
                     if valor is not None and not callable(valor))


def _detalle_alertas(argumentos: Dict[str, Any], resultado: Any) -> str:
    """Detalles de una operación que genera el CSV de alertas."""
    return (f"diferencia={argumentos.get('diferencia')}, dólar {argumentos.get('tipo')}: "
            f"{resultado.comparados} libros comparados, {resultado.alertas} alertas, "
            f"archivo {os.path.basename(resultado.ruta)}")


def _detalle_reporte(argumentos: Dict[str, Any], resultado: Any) -> str:
    """Detalles de la generación del reporte en Excel."""
    return (f"dólar {argumentos.get('tipo')}: {resultado.libros} libros, "
            f"{resultado.sin_coincidencia} sin coincidencia, "
            f"archivo {os.path.basename(resultado.ruta)}")


def _registrar_operacion(firma: inspect.Signature, accion: str,
                         detalle: Optional[Callable[[Dict[str, Any], Any], str]],
                         args: tuple, kwargs: dict, resultado: Any,
                         error: Optional[Exception]) -> None:
    """Guarda en la auditoría una operación terminada, sin interrumpir nunca a la operación."""
    servicio = _servicio_auditoria
    if servicio is None:
        return
    try:
        argumentos = dict(firma.bind(*args, **kwargs).arguments)
        objeto = getattr(argumentos.pop("self", None), "ENTIDAD", "")
        nombre = accion.format_map(_SinFaltantes(argumentos, objeto=objeto)).strip()
        if error is not None:
            partes = [f"ERROR ({type(error).__name__}): {error}"]
        else:
            partes = []
        if error is None and detalle is not None:
            partes.append(detalle(argumentos, resultado))
        else:
            partes.append(_describir_argumentos(argumentos))
            if error is None and resultado is not None and not isinstance(resultado, bool):
                partes.append(f"resultado: {_resumir(resultado)}")
        detalles = " | ".join(parte for parte in partes if parte)
        servicio.registrar(nombre, detalles)
    except Exception as fallo:  # la auditoría nunca debe romper la operación que audita
        logging.getLogger(__name__).warning("No se pudo registrar la auditoría: %s", fallo)


def auditar(accion: str, detalle: Optional[Callable[[Dict[str, Any], Any], str]] = None):
    """Decorador que registra en la auditoría cada vez que se ejecuta una operación.

    Guarda la acción, la fecha y los detalles. Si la operación falla, también queda el
    registro (con el error) y la excepción se vuelve a lanzar. Si una operación auditada
    llama a otra, solo se registra la de más afuera, para no duplicar.

    Args:
        accion: nombre de la operación. Puede incluir {objeto} (el valor de ENTIDAD del
            servicio, por ejemplo "libro") y el nombre de cualquier argumento, como {titulo}.
        detalle: función opcional (argumentos, resultado) -> texto para armar los detalles.
            Por defecto se guardan los argumentos y el resultado.
    """
    def decorador(funcion):
        """Envuelve la función para que cada ejecución quede registrada."""
        firma = inspect.signature(funcion)

        @functools.wraps(funcion)
        def envoltura(*args, **kwargs):
            """Ejecuta la función original y registra el resultado o el error."""
            nivel = getattr(_contexto, "nivel", 0)
            _contexto.nivel = nivel + 1
            try:
                resultado = funcion(*args, **kwargs)
            except Exception as error:
                if nivel == 0:
                    _registrar_operacion(firma, accion, detalle, args, kwargs, None, error)
                raise
            else:
                if nivel == 0:
                    _registrar_operacion(firma, accion, detalle, args, kwargs, resultado, None)
                return resultado
            finally:
                _contexto.nivel = nivel
        return envoltura
    return decorador


class ServicioAuditoria:
    """Registra y consulta el historial de operaciones del sistema."""

    def __init__(self, repositorio: RepositorioAuditoria) -> None:
        """Guarda el repositorio donde se escribe la tabla de auditoría."""
        self._repositorio = repositorio

    @staticmethod
    def _ahora() -> datetime.datetime:
        """Fecha y hora actuales en Argentina (UTC-3), sin segundos fraccionarios."""
        zona = ServicioCotizacionDolar._ZONA_ARGENTINA
        return datetime.datetime.now(zona).replace(tzinfo=None, microsecond=0)

    def registrar(self, accion: str, detalles: str = "") -> Auditoria:
        """Guarda una operación en la tabla de auditoría con la fecha actual.

        Args:
            accion: nombre de la operación (se corta a 120 caracteres).
            detalles: datos de la operación y su resultado.
        """
        registro = Auditoria(self._repositorio.siguiente_id(), accion[:120], self._ahora(),
                             detalles)
        return self._repositorio.crear(registro)

    @auditar("Ver historial de auditoría",
             detalle=lambda argumentos, resultado: f"{len(resultado)} registros")
    def historial(self, limite: Optional[int] = None) -> List[Auditoria]:
        """Devuelve las operaciones registradas, de la más reciente a la más antigua."""
        return self._repositorio.leer_recientes(limite)


class FilaBimonetaria(NamedTuple):
    """Un libro con su precio en pesos y en otra moneda (None si no tiene)."""

    libro: Libro
    pesos: Optional[Decimal]
    otra_moneda: Optional[Decimal]


class ServicioBase(Generic[T]):
    """CRUD común a los servicios de entidades con id propio."""

    ENTIDAD = "registro"

    def __init__(self, repositorio: RepositorioDB[T]) -> None:
        """Guarda el repositorio con el que se leen y escriben las entidades."""
        self._repositorio = repositorio

    def listar(self) -> List[T]:
        """Devuelve todas las entidades."""
        return self._repositorio.leer_todos()

    def _obtener_por_id(self, id: int) -> T:
        """Uso interno: devuelve la entidad o lanza ValueError si no existe."""
        entidad = self._repositorio.leer_por_id(id)
        if entidad is None:
            raise ValueError(f"No existe la entidad con id {id}.")
        return entidad

    @auditar("Modificar {objeto}")
    def modificar(self, entidad: T) -> T:
        """Persiste los cambios de una entidad existente."""
        self._obtener_por_id(entidad.id)
        return self._repositorio.actualizar(entidad)

    @auditar("Eliminar {objeto}")
    def eliminar(self, entidad: T) -> None:
        """Elimina la entidad si ninguna otra la está usando."""
        self._obtener_por_id(entidad.id)
        self._validar_eliminacion(entidad)
        self._repositorio.eliminar(entidad.id)

    def _validar_eliminacion(self, entidad: T) -> None:
        """Punto de extensión: lanza ValueError si la entidad no puede borrarse."""

    def _validar_nombre_unico(self, nombre: str, id_actual: Optional[int] = None) -> None:
        """Evita nombres repetidos (sin distinguir mayúsculas)."""
        for existente in self._repositorio.leer_todos():
            if existente.nombre.lower() == nombre.strip().lower() and existente.id != id_actual:
                raise ValueError(f"Ya existe '{nombre}'.")


class ServicioGenero(ServicioBase[Genero]):
    """Lógica de géneros literarios."""

    ENTIDAD = "género"

    def __init__(self, repositorio: RepositorioGenero, repo_libro: RepositorioLibro) -> None:
        """Guarda el repositorio de géneros y el de libros (para validar las bajas)."""
        super().__init__(repositorio)
        self._repo_libro = repo_libro

    @auditar("Crear {objeto}")
    def crear(self, nombre: str) -> Genero:
        """Crea un género nuevo; no se permiten nombres repetidos."""
        self._validar_nombre_unico(nombre)
        genero = Genero(self._repositorio.siguiente_id(), nombre)
        return self._repositorio.crear(genero)

    @auditar("Modificar {objeto}")
    def modificar(self, genero: Genero) -> Genero:
        """Guarda los cambios de un género, sin permitir un nombre ya usado por otro."""
        self._validar_nombre_unico(genero.nombre, genero.id)
        return super().modificar(genero)

    def _validar_eliminacion(self, entidad: Genero) -> None:
        """Impide eliminar un género si hay libros que lo usan."""
        if any(libro.genero == entidad for libro in self._repo_libro.leer_todos()):
            raise ValueError("No se puede eliminar: hay libros con este género.")


class ServicioEditorial(ServicioBase[Editorial]):
    """Lógica de editoriales."""

    ENTIDAD = "editorial"

    def __init__(self, repositorio: RepositorioEditorial, repo_libro: RepositorioLibro) -> None:
        """Guarda el repositorio de editoriales y el de libros (para validar las bajas)."""
        super().__init__(repositorio)
        self._repo_libro = repo_libro

    @auditar("Crear {objeto}")
    def crear(self, nombre: str, pais: str = "Argentina") -> Editorial:
        """Crea una editorial nueva (país por defecto: Argentina), sin nombres repetidos."""
        self._validar_nombre_unico(nombre)
        editorial = Editorial(self._repositorio.siguiente_id(), nombre, pais)
        return self._repositorio.crear(editorial)

    @auditar("Modificar {objeto}")
    def modificar(self, editorial: Editorial) -> Editorial:
        """Guarda los cambios de una editorial, sin permitir un nombre ya usado por otra."""
        self._validar_nombre_unico(editorial.nombre, editorial.id)
        return super().modificar(editorial)

    def _validar_eliminacion(self, entidad: Editorial) -> None:
        """Impide eliminar una editorial si tiene libros."""
        if any(libro.editorial == entidad for libro in self._repo_libro.leer_todos()):
            raise ValueError("No se puede eliminar: hay libros de esta editorial.")


class ServicioMoneda(ServicioBase[Moneda]):
    """Lógica de monedas."""

    ENTIDAD = "moneda"

    def __init__(self, repositorio: RepositorioMoneda, repo_precio: RepositorioPrecio) -> None:
        """Guarda el repositorio de monedas y el de precios (para validar las bajas)."""
        super().__init__(repositorio)
        self._repo_precio = repo_precio

    @auditar("Crear {objeto}")
    def crear(self, codigo: str, nombre: str, simbolo: str) -> Moneda:
        """Crea una moneda nueva; no se permiten códigos repetidos."""
        for existente in self._repositorio.leer_todos():
            if existente.codigo == codigo.strip().upper():
                raise ValueError(f"Ya existe la moneda {codigo}.")
        moneda = Moneda(self._repositorio.siguiente_id(), codigo, nombre, simbolo)
        return self._repositorio.crear(moneda)

    def obtener_por_codigo(self, codigo: str) -> Moneda:
        """Busca una moneda por su código (ARS, USD...)."""
        for moneda in self._repositorio.leer_todos():
            if moneda.codigo == codigo.upper():
                return moneda
        raise ValueError(f"No existe la moneda {codigo}.")

    def _validar_eliminacion(self, entidad: Moneda) -> None:
        """Impide eliminar una moneda si hay precios expresados en ella."""
        if any(precio.moneda == entidad for precio in self._repo_precio.leer_todos()):
            raise ValueError("No se puede eliminar: hay precios en esta moneda.")


class ServicioTipoCotizacion(ServicioBase[TipoCotizacion]):
    """Lógica de tipos de cotización del dólar."""

    ENTIDAD = "tipo de cotización"

    def __init__(self, repositorio: RepositorioTipoCotizacion,
                 repo_cotizacion: RepositorioCotizacionDolar) -> None:
        """Guarda el repositorio de tipos y el de cotizaciones (para validar las bajas)."""
        super().__init__(repositorio)
        self._repo_cotizacion = repo_cotizacion

    @auditar("Crear {objeto}")
    def crear(self, nombre: str) -> TipoCotizacion:
        """Crea un tipo de cotización nuevo; no se permiten nombres repetidos."""
        self._validar_nombre_unico(nombre)
        tipo = TipoCotizacion(self._repositorio.siguiente_id(), nombre)
        return self._repositorio.crear(tipo)

    @auditar("Modificar {objeto}")
    def modificar(self, tipo: TipoCotizacion) -> TipoCotizacion:
        """Guarda los cambios de un tipo, sin permitir un nombre ya usado por otro."""
        self._validar_nombre_unico(tipo.nombre, tipo.id)
        return super().modificar(tipo)

    def _validar_eliminacion(self, entidad: TipoCotizacion) -> None:
        """Impide eliminar un tipo de cotización si tiene cotizaciones registradas."""
        if self._repo_cotizacion.leer_historico_por_tipo(entidad.id):
            raise ValueError("No se puede eliminar: tiene cotizaciones registradas.")


class ServicioLibro(ServicioBase[Libro]):
    """Lógica de libros."""

    ENTIDAD = "libro"

    def __init__(self, repositorio: RepositorioLibro, repo_precio: RepositorioPrecio,
                 repo_stock: RepositorioStock) -> None:
        """Guarda los repositorios de libros, de precios y de stock."""
        super().__init__(repositorio)
        self._repo_precio = repo_precio
        self._repo_stock = repo_stock

    def _validar_isbn_unico(self, isbn: str, id_actual: Optional[int] = None) -> None:
        """Verifica que no exista otro libro con el mismo ISBN (los guiones no cuentan)."""
        limpio = isbn.replace("-", "").strip()
        for existente in self._repositorio.leer_todos():
            if existente.isbn == limpio and existente.id != id_actual:
                raise ValueError(f"Ya existe un libro con el ISBN {isbn}.")

    @auditar("Crear {objeto}")
    def crear(self, isbn: str, titulo: str, autor: str,
              editorial: Editorial, genero: Genero) -> Libro:
        """Crea un libro nuevo; no se permiten ISBN repetidos."""
        self._validar_isbn_unico(isbn)
        libro = Libro(self._repositorio.siguiente_id(), isbn, titulo, autor, editorial, genero)
        return self._repositorio.crear(libro)

    @auditar("Modificar {objeto}")
    def modificar(self, libro: Libro) -> Libro:
        """Guarda los cambios de un libro, sin permitir un ISBN ya usado por otro."""
        self._validar_isbn_unico(libro.isbn, libro.id)
        return super().modificar(libro)

    @auditar("Buscar libros")
    def buscar(self, texto: str) -> List[Libro]:
        """Busca libros por título o autor (sin distinguir mayúsculas)."""
        texto = texto.strip().lower()
        return [libro for libro in self._repositorio.leer_todos()
                if texto in libro.titulo.lower() or texto in libro.autor.lower()]

    @auditar("Eliminar {objeto}")
    def eliminar(self, libro: Libro) -> None:
        """Elimina el libro junto con sus precios y su stock."""
        self._obtener_por_id(libro.id)
        for precio in self._repo_precio.leer_todos():
            if precio.libro == libro:
                self._repo_precio.eliminar(precio.id)
        self._repo_stock.eliminar(libro.id)
        self._repositorio.eliminar(libro.id)


class IProveedorCotizacion(abc.ABC):
    """Fuente externa de cotizaciones del dólar."""

    @abc.abstractmethod
    def obtener_todas(self) -> List[Dict[str, object]]:
        """Devuelve todas las cotizaciones disponibles.

        Cada una es un diccionario con las claves 'casa', 'nombre', 'compra', 'venta' y
        'fechaActualizacion'.
        """


class DolarApiProveedor(IProveedorCotizacion):
    """Obtiene las cotizaciones en tiempo real desde la API configurada en el .env."""

    def __init__(self, url: Optional[str] = None) -> None:
        """Args:
            url: dirección de la API. Si no se indica, se lee API_URL del .env.
        """
        self._url = url

    def _direccion(self) -> str:
        """Devuelve la URL de la API: la indicada al crear el proveedor o API_URL del entorno."""
        cargar_variables_entorno()
        url = self._url or os.getenv("API_URL")
        if not url:
            raise ValueError("Falta API_URL: definila en el archivo .env.")
        return url

    def obtener_todas(self) -> List[Dict[str, object]]:
        """Consulta la API y devuelve la lista de cotizaciones tal como la entrega."""
        respuesta = requests.get(self._direccion(), timeout=10)
        respuesta.raise_for_status()
        datos = respuesta.json()
        if not isinstance(datos, list):
            raise ValueError("La API devolvió un formato inesperado.")
        return datos


class ServicioCotizacionDolar:
    """Lógica de cotizaciones del dólar."""

    ENTIDAD = "cotización del dólar"

    # Identificador que usa la API ("casa") -> nombre del tipo de cotización en el sistema.
    TIPOS_POR_CASA = {
        "oficial": "Oficial", "blue": "Blue", "bolsa": "MEP", "contadoconliqui": "CCL",
        "mayorista": "Mayorista", "tarjeta": "Tarjeta", "cripto": "Cripto"}

    # Argentina no tiene horario de verano: la hora local es siempre UTC-3.
    _ZONA_ARGENTINA = datetime.timezone(datetime.timedelta(hours=-3))

    def __init__(self, repositorio: RepositorioCotizacionDolar,
                 repo_tipo: RepositorioTipoCotizacion,
                 proveedor: IProveedorCotizacion) -> None:
        """Guarda los repositorios de cotizaciones y tipos, y el proveedor que consulta la API."""
        self._repositorio = repositorio
        self._repo_tipo = repo_tipo
        self._proveedor = proveedor

    @auditar("Registrar {objeto}")
    def registrar(self, tipo: TipoCotizacion, fecha: datetime.date,
                  valor_compra, valor_venta) -> CotizacionDolar:
        """Registra manualmente una cotización."""
        cotizacion = CotizacionDolar(tipo, fecha, valor_compra, valor_venta)
        return self._repositorio.crear(cotizacion)

    @auditar("Modificar {objeto}")
    def modificar(self, cotizacion: CotizacionDolar) -> CotizacionDolar:
        """Guarda los cambios de una cotización."""
        return self._repositorio.actualizar(cotizacion)

    @auditar("Eliminar {objeto}")
    def eliminar(self, tipo: TipoCotizacion, fecha: datetime.date) -> bool:
        """Elimina la cotización de un tipo en una fecha. Devuelve True si se eliminó."""
        return self._repositorio.eliminar(tipo.id, fecha)

    def historico(self, tipo: TipoCotizacion) -> List[CotizacionDolar]:
        """Cotizaciones de un tipo, de la más antigua a la más reciente."""
        return self._repositorio.leer_historico_por_tipo(tipo.id)

    def listar(self) -> List[CotizacionDolar]:
        """Devuelve todas las cotizaciones registradas."""
        return self._repositorio.leer_todos()

    def ultima(self, tipo: TipoCotizacion) -> Optional[CotizacionDolar]:
        """Cotización más reciente de un tipo, o None si no hay."""
        historico = self.historico(tipo)
        return historico[-1] if historico else None

    @classmethod
    def _fecha_de(cls, texto: object) -> datetime.date:
        """Fecha local (Argentina) de la actualización. Si no se puede leer, la de hoy."""
        try:
            momento = datetime.datetime.fromisoformat(str(texto).replace("Z", "+00:00"))
        except ValueError:
            return datetime.date.today()
        if momento.tzinfo is None:
            return momento.date()
        return momento.astimezone(cls._ZONA_ARGENTINA).date()

    @classmethod
    def _interpretar(cls, dato: Dict[str, object]) -> Tuple[str, datetime.date, Decimal, Decimal]:
        """Convierte un elemento de la respuesta en (tipo, fecha, compra, venta).

        Si la API no informa el valor de compra, se toma el de venta.
        """
        try:
            casa = str(dato["casa"]).lower()
            nombre = cls.TIPOS_POR_CASA.get(casa) or str(dato.get("nombre") or casa)
            venta = Decimal(str(dato["venta"]))
            compra = venta if dato.get("compra") is None else Decimal(str(dato["compra"]))
        except (KeyError, TypeError, InvalidOperation) as error:
            raise ValueError(
                f"La API devolvió una cotización con formato inesperado: {dato}") from error
        if compra < 0 or venta < 0:
            raise ValueError(f"La API devolvió una cotización con valores negativos: {dato}")
        return nombre, cls._fecha_de(dato.get("fechaActualizacion")), compra, venta

    @auditar("Obtener cotizaciones por API")
    def obtener_cotizaciones(self) -> List[CotizacionDolar]:
        """Busca las cotizaciones del dólar en la API, las descarga y las registra en la base.

        Por cada cotización recibida: si ya hay una para ese tipo y fecha, se actualiza; si
        no, se crea. Si el tipo de cotización no existe, también se crea. Se puede repetir
        sin duplicar registros. Los datos se validan todos antes de escribir en la base.

        Returns:
            Las cotizaciones registradas.

        Raises:
            requests.RequestException: si falla la conexión con la API.
            ValueError: si falta API_URL o la respuesta no es válida.
        """
        datos = [self._interpretar(dato) for dato in self._proveedor.obtener_todas()]
        if not datos:
            raise ValueError("La API no devolvió cotizaciones.")
        tipos = {tipo.nombre.lower(): tipo for tipo in self._repo_tipo.leer_todos()}
        registradas = []
        for nombre, fecha, compra, venta in datos:
            tipo = tipos.get(nombre.lower())
            if tipo is None:
                tipo = self._repo_tipo.crear(
                    TipoCotizacion(self._repo_tipo.siguiente_id(), nombre))
                tipos[nombre.lower()] = tipo
            existente = self._repositorio.leer_por_tipo_y_fecha(tipo.id, fecha)
            if existente is None:
                registradas.append(
                    self._repositorio.crear(CotizacionDolar(tipo, fecha, compra, venta)))
            else:
                existente.valor_compra = compra
                existente.valor_venta = venta
                registradas.append(self._repositorio.actualizar(existente))
        return registradas


class ServicioPrecio(ServicioBase[Precio]):
    """Lógica de precios y cotización de libros."""

    ENTIDAD = "precio"

    def __init__(self, repositorio: RepositorioPrecio,
                 servicio_cotizacion: ServicioCotizacionDolar) -> None:
        """Guarda el repositorio de precios y el servicio de cotizaciones del dólar."""
        super().__init__(repositorio)
        self._servicio_cotizacion = servicio_cotizacion

    def precios_de_libro(self, libro: Libro) -> List[Precio]:
        """Todos los precios cargados para un libro."""
        return [precio for precio in self._repositorio.leer_todos() if precio.libro == libro]

    @auditar("Crear {objeto}")
    def crear(self, libro: Libro, moneda: Moneda, valor) -> Precio:
        """Crea un precio para un libro; no puede tener dos precios en la misma moneda."""
        if any(precio.moneda == moneda for precio in self.precios_de_libro(libro)):
            raise ValueError(f"El libro ya tiene un precio en {moneda.codigo}.")
        precio = Precio(self._repositorio.siguiente_id(), libro, moneda, valor)
        return self._repositorio.crear(precio)

    @staticmethod
    def _pesos_de(libro: Libro, precios: List[Precio],
                  cotizacion: Optional[CotizacionDolar], tipo: TipoCotizacion) -> Decimal:
        """Precio en ARS de un libro a partir de sus precios y de la cotización del dólar."""
        en_usd = next((p for p in precios if p.moneda.codigo == CODIGO_USD), None)
        if en_usd is not None:
            if cotizacion is None:
                raise ValueError(f"No hay cotización registrada para {tipo.nombre}.")
            pesos = en_usd.valor * cotizacion.valor_venta
            return pesos.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        en_ars = next((p for p in precios if p.moneda.codigo == CODIGO_ARS), None)
        if en_ars is not None:
            return en_ars.valor
        raise ValueError(f"El libro '{libro.titulo}' no tiene precio en USD ni en ARS.")

    def cotizar_en_pesos(self, libro: Libro, tipo: TipoCotizacion) -> Decimal:
        """Precio del libro en ARS. Si está en USD lo convierte con la última venta del tipo."""
        return self._pesos_de(libro, self.precios_de_libro(libro),
                              self._servicio_cotizacion.ultima(tipo), tipo)

    def precios_en_pesos(self, tipo: TipoCotizacion) -> Dict[int, Decimal]:
        """Precio en pesos de cada libro que tiene precio, indexado por el id del libro.

        Usa la misma conversión que la lista bimonetaria: si el libro tiene precio en USD
        se convierte con la última venta del tipo de dólar, y si no, se usa el de ARS. Los
        libros sin precio en USD ni ARS no aparecen en el resultado.

        Raises:
            ValueError: si hay precios en dólares y no hay cotización registrada del tipo.
        """
        precios = self._repositorio.leer_todos()
        cotizacion = self._servicio_cotizacion.ultima(tipo)
        if cotizacion is None and any(p.moneda.codigo == CODIGO_USD for p in precios):
            raise ValueError(f"No hay cotización registrada para {tipo.nombre}. "
                             "Obtenelas por API o cargalas a mano.")
        por_libro: Dict[int, List[Precio]] = {}
        libros: Dict[int, Libro] = {}
        for precio in precios:
            por_libro.setdefault(precio.libro.id, []).append(precio)
            libros[precio.libro.id] = precio.libro
        resultado: Dict[int, Decimal] = {}
        for id_libro, libro in libros.items():
            try:
                resultado[id_libro] = self._pesos_de(libro, por_libro[id_libro], cotizacion, tipo)
            except ValueError:
                continue
        return resultado

    def monedas_con_precio(self, excluir_codigo: Optional[str] = None) -> List[Moneda]:
        """Monedas en las que hay al menos un precio cargado, ordenadas por id."""
        monedas = {precio.moneda for precio in self._repositorio.leer_todos()}
        return sorted((m for m in monedas if m.codigo != excluir_codigo), key=lambda m: m.id)

    @auditar("Generar lista bimonetaria")
    def lista_bimonetaria(self, tipo: TipoCotizacion, moneda: Moneda) -> List[FilaBimonetaria]:
        """Lista de precios en pesos y en otra moneda para todos los libros con precio.

        El precio en pesos se calcula con la última cotización del tipo de dólar indicado.
        El precio en la otra moneda es el que está cargado: no se convierte, porque las
        cotizaciones disponibles son solo del dólar.

        Raises:
            ValueError: si hay precios en dólares y no hay cotización registrada del tipo.
        """
        precios = self._repositorio.leer_todos()
        cotizacion = self._servicio_cotizacion.ultima(tipo)
        if cotizacion is None and any(p.moneda.codigo == CODIGO_USD for p in precios):
            raise ValueError(f"No hay cotización registrada para {tipo.nombre}. "
                             "Obtenelas por API o cargalas a mano.")
        por_libro: Dict[int, List[Precio]] = {}
        libros: Dict[int, Libro] = {}
        for precio in precios:
            por_libro.setdefault(precio.libro.id, []).append(precio)
            libros[precio.libro.id] = precio.libro
        filas = []
        for id_libro in sorted(libros):
            libro, del_libro = libros[id_libro], por_libro[id_libro]
            try:
                pesos: Optional[Decimal] = self._pesos_de(libro, del_libro, cotizacion, tipo)
            except ValueError:
                pesos = None
            otra = next((p.valor for p in del_libro if p.moneda == moneda), None)
            filas.append(FilaBimonetaria(libro, pesos, otra))
        return filas

    @auditar("Exportar precios a CSV")
    def exportar_precios_csv(self, nombre_archivo: str = "precios.csv",
                             carpeta: str = CARPETA_EXPORTACION) -> Tuple[str, int]:
        """Exporta los precios cargados a un CSV: una fila por libro y una columna por moneda.

        Las columnas de moneda cubren todas las monedas que tienen algún precio; queda
        vacía la celda de un libro que no tiene precio en esa moneda.

        Returns:
            La ruta completa del archivo y la cantidad de libros exportados.
        """
        nombre = os.path.basename(nombre_archivo.strip())
        if not nombre:
            raise ValueError("El nombre del archivo no puede estar vacío.")
        if not nombre.lower().endswith(".csv"):
            nombre += ".csv"
        precios = self._repositorio.leer_todos()
        if not precios:
            raise ValueError("No hay precios cargados para exportar.")
        monedas = sorted({precio.moneda for precio in precios}, key=lambda m: m.id)
        libros: Dict[int, Libro] = {}
        valores: Dict[Tuple[int, int], Decimal] = {}
        for precio in precios:
            libros[precio.libro.id] = precio.libro
            valores[(precio.libro.id, precio.moneda.id)] = precio.valor
        os.makedirs(carpeta, exist_ok=True)
        ruta = os.path.join(carpeta, nombre)
        # utf-8 con BOM: así Excel muestra bien los acentos al abrir el archivo.
        with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(["isbn", "titulo", "autor"] + [f"precio_{m.codigo}" for m in monedas])
            for id_libro in sorted(libros):
                libro = libros[id_libro]
                escritor.writerow([libro.isbn, libro.titulo, libro.autor] + [
                    valores.get((id_libro, m.id), "") for m in monedas])
        return ruta, len(libros)


class ServicioStock:
    """Lógica de stock, un registro por libro."""

    ENTIDAD = "stock"

    def __init__(self, repositorio: RepositorioStock) -> None:
        """Guarda el repositorio de stock."""
        self._repositorio = repositorio

    @auditar("Crear {objeto}")
    def crear(self, libro: Libro, cantidad: int) -> Stock:
        """Crea el registro de stock de un libro con la cantidad indicada."""
        return self._repositorio.crear(Stock(libro, cantidad))

    def listar(self) -> List[Stock]:
        """Devuelve el stock de todos los libros."""
        return self._repositorio.leer_todos()

    def obtener(self, libro: Libro) -> Stock:
        """Devuelve el stock de un libro; falla si no tiene stock cargado."""
        stock = self._repositorio.leer_por_libro(libro.id)
        if stock is None:
            raise ValueError(f"El libro '{libro.titulo}' no tiene stock cargado.")
        return stock

    @auditar("Modificar {objeto}")
    def modificar(self, stock: Stock) -> Stock:
        """Guarda los cambios de un registro de stock."""
        return self._repositorio.actualizar(stock)

    @auditar("Ajustar stock")
    def ajustar(self, libro: Libro, delta: int) -> Stock:
        """Suma o resta unidades; no permite quedar con stock negativo."""
        stock = self.obtener(libro)
        if stock.cantidad + delta < 0:
            raise ValueError(f"Stock insuficiente: hay {stock.cantidad} unidades.")
        stock.cantidad = stock.cantidad + delta
        return self._repositorio.actualizar(stock)

    @auditar("Eliminar {objeto}")
    def eliminar(self, libro: Libro) -> bool:
        """Elimina el stock de un libro. Devuelve True si se eliminó."""
        return self._repositorio.eliminar(libro.id)

    @auditar("Consultar stock bajo")
    def con_stock_bajo(self, umbral: int = 5) -> List[Stock]:
        """Libros con una cantidad menor o igual al umbral."""
        return [stock for stock in self._repositorio.leer_todos() if stock.cantidad <= umbral]


# ---------------------------------------------------------------- scraping y alertas
# Palabras que no ayudan a reconocer un título ("El Aleph" también se publica como "Aleph").
PALABRAS_VACIAS = frozenset({"el", "la", "los", "las", "un", "una", "unos", "unas", "de",
                             "del", "al", "y", "e", "en", "a", "o", "u"})


def _palabras(texto: str) -> List[str]:
    """Pasa un texto a minúsculas, sin tildes ni signos, y lo separa en palabras."""
    sin_tildes = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.findall(r"[a-z0-9]+", sin_tildes.lower())


def _significativas(titulo: str) -> frozenset:
    """Palabras que identifican un título: sin tildes, mayúsculas, artículos ni aclaraciones.

    Lo que va entre paréntesis o corchetes (edición, tipo de tapa, etc.) no cuenta.
    """
    palabras = frozenset(_palabras(re.sub(r"[(\[][^)\]]*[)\]]", " ", titulo)))
    return frozenset(palabra for palabra in palabras if palabra not in PALABRAS_VACIAS) or palabras


def elegir_coincidencia(libro: Libro, candidatos: Sequence[PrecioWeb]) -> Optional[PrecioWeb]:
    """Elige, entre los resultados de la búsqueda, la publicación que corresponde al libro.

    El título de la publicación tiene que ser el mismo que el del libro: se ignoran las
    tildes, las mayúsculas, los artículos y preposiciones, y lo que va entre paréntesis
    (por ejemplo "(TAPA DURA)"). Un título con palabras de más, como "FICCIONES CON CIENCIA"
    para el libro "Ficciones", es otra obra y no cuenta. Si hay varias ediciones que
    cumplen, se elige la que el sitio mostró primero. No se usa el ISBN porque cambia
    con cada edición.

    Returns:
        La publicación elegida, o None si ninguna coincide.
    """
    buscadas = _significativas(libro.titulo)
    validos = [candidato for candidato in candidatos
               if _significativas(candidato.titulo_web) == buscadas]
    return min(validos, key=lambda candidato: candidato.posicion) if validos else None


class FilaComparacion(NamedTuple):
    """Un libro con su precio interno y el de la publicación web elegida (None si falta)."""

    libro: Libro
    precio_interno: Optional[Decimal]
    precio_web: Optional[Decimal]
    diferencia: Optional[Decimal]
    titulo_web: str
    url: str
    fecha_extraccion: Optional[datetime.datetime]


class ResultadoAlertas(NamedTuple):
    """Resumen de una generación de alertas."""

    ruta: str
    comparados: int
    sin_coincidencia: int
    alertas: int
    fecha_extraccion: datetime.datetime


class ResultadoReporte(NamedTuple):
    """Resumen de la generación del reporte en Excel."""

    ruta: str
    libros: int
    sin_coincidencia: int
    fecha_extraccion: datetime.datetime


class ServicioScraping:
    """Ejecuta el scraper, compara los precios internos con los de la web y genera alertas."""

    COLUMNAS_REPORTE = ["Libro", "Precio interno", "Precio web", "Diferencia",
                        "Fecha de extracción"]
    COLUMNAS_ALERTAS = ["libro", "precio_interno", "precio_web", "diferencia",
                        "diferencia_porcentual", "alerta", "titulo_web", "url",
                        "fecha_extraccion"]

    def __init__(self, repo_libro: RepositorioLibro, repo_precio_web: RepositorioPrecioWeb,
                 servicio_precio: ServicioPrecio, lanzador: Callable[..., int]) -> None:
        """Guarda las dependencias.

        Args:
            repo_libro: repositorio del catálogo.
            repo_precio_web: repositorio de los precios obtenidos por el scraper.
            servicio_precio: calcula el precio interno en pesos.
            lanzador: función que corre el scraper y devuelve su código de salida
                (0 si terminó bien). Recibe fecha_extraccion como argumento con nombre.
        """
        self._repo_libro = repo_libro
        self._repo_precio_web = repo_precio_web
        self._servicio_precio = servicio_precio
        self._lanzador = lanzador

    @staticmethod
    def _ahora() -> datetime.datetime:
        """Fecha y hora actuales en Argentina (UTC-3), sin segundos fraccionarios."""
        zona = ServicioCotizacionDolar._ZONA_ARGENTINA
        return datetime.datetime.now(zona).replace(tzinfo=None, microsecond=0)

    def comparar(self, tipo: TipoCotizacion,
                 fecha_extraccion: Optional[datetime.datetime] = None) -> List[FilaComparacion]:
        """Compara el precio interno (en pesos) de cada libro con el de su publicación web.

        Args:
            tipo: tipo de dólar para convertir los precios en dólares.
            fecha_extraccion: corrida del scraper a usar. Por defecto, la más reciente.

        Returns:
            Una fila por libro del catálogo, ordenadas por id. Si el libro no tiene precio
            interno o no hay una publicación que coincida, esos campos quedan en None.

        Raises:
            ValueError: si hay precios en dólares y no hay cotización del tipo indicado.
        """
        return self._comparar(tipo, self._repo_precio_web.leer_por_extraccion(fecha_extraccion))

    def _comparar(self, tipo: TipoCotizacion,
                  precios_web: List[PrecioWeb]) -> List[FilaComparacion]:
        """Arma la comparación a partir de los precios web de una corrida ya leídos."""
        internos = self._servicio_precio.precios_en_pesos(tipo)
        momento = precios_web[0].fecha_extraccion if precios_web else None
        por_libro: Dict[int, List[PrecioWeb]] = {}
        for precio_web in precios_web:
            por_libro.setdefault(precio_web.libro.id, []).append(precio_web)
        filas = []
        for libro in self._repo_libro.leer_todos():
            elegido = elegir_coincidencia(libro, por_libro.get(libro.id, []))
            interno = internos.get(libro.id)
            web = None if elegido is None else elegido.precio
            diferencia = None if interno is None or web is None else web - interno
            filas.append(FilaComparacion(
                libro, interno, web, diferencia,
                "" if elegido is None else elegido.titulo_web,
                "" if elegido is None else elegido.url,
                momento))
        return filas

    def _leer_corrida(self, fecha_extraccion: Optional[datetime.datetime]) -> List[PrecioWeb]:
        """Lee los precios web de una corrida (por defecto la más reciente).

        Raises:
            ValueError: si no hay precios guardados de esa corrida.
        """
        precios_web = self._repo_precio_web.leer_por_extraccion(fecha_extraccion)
        if not precios_web:
            if fecha_extraccion is None:
                raise ValueError("No hay precios web guardados: ejecutá el scraping primero.")
            raise ValueError("El scraper no guardó precios. Revisá la conexión o si el sitio "
                             "bloquea las búsquedas.")
        return precios_web

    @auditar("Generar alertas", detalle=_detalle_alertas)
    def generar_alertas(self, diferencia: Decimal, tipo: TipoCotizacion,
                        carpeta: str = CARPETA_EXPORTACION,
                        fecha_extraccion: Optional[datetime.datetime] = None
                        ) -> ResultadoAlertas:
        """Genera el CSV de alertas: los libros cuya diferencia de precio llega al umbral.

        El archivo se crea siempre, aunque no haya alertas (queda solo con los títulos de
        las columnas). La diferencia es precio web menos precio interno, en pesos.

        Args:
            diferencia: diferencia mínima (en pesos, sin signo) que dispara una alerta.
            tipo: tipo de dólar para convertir los precios internos en dólares.
            carpeta: carpeta donde se guarda el archivo.
            fecha_extraccion: corrida del scraper a usar. Por defecto, la más reciente.

        Raises:
            ValueError: si la diferencia es negativa, si no hay precios web de esa corrida
                o si falta la cotización del dólar.
        """
        umbral = Decimal(str(diferencia))
        if umbral < 0:
            raise ValueError("La diferencia de precios no puede ser negativa.")
        precios_web = self._leer_corrida(fecha_extraccion)
        momento = precios_web[0].fecha_extraccion
        filas = self._comparar(tipo, precios_web)
        comparados = [f for f in filas if f.diferencia is not None]
        alertas = [f for f in comparados if abs(f.diferencia) >= umbral]
        os.makedirs(carpeta, exist_ok=True)
        ruta = os.path.join(carpeta, f"alertas_{momento:%Y%m%d_%H%M%S}.csv")
        # utf-8 con BOM: así Excel muestra bien los acentos al abrir el archivo.
        with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(self.COLUMNAS_ALERTAS)
            for fila in alertas:
                porcentual = (fila.diferencia / fila.precio_interno * 100).quantize(
                    Decimal("0.01"), rounding=ROUND_HALF_UP) if fila.precio_interno else ""
                escritor.writerow([
                    fila.libro.titulo, fila.precio_interno, fila.precio_web,
                    fila.diferencia, porcentual,
                    "web más caro" if fila.diferencia > 0 else "web más barato",
                    fila.titulo_web, fila.url, f"{fila.fecha_extraccion:%Y-%m-%d %H:%M:%S}"])
        sin_coincidencia = sum(1 for f in filas if f.precio_web is None)
        return ResultadoAlertas(ruta, len(comparados), sin_coincidencia, len(alertas), momento)

    @auditar("Ejecutar scraping", detalle=_detalle_alertas)
    def ejecutar(self, diferencia: Decimal, tipo: TipoCotizacion,
                 carpeta: str = CARPETA_EXPORTACION) -> ResultadoAlertas:
        """Corre el scraper y, al terminar, genera el archivo de alertas de esa corrida.

        Las validaciones (umbral y cotización del dólar) se hacen antes de iniciar el
        scraping, para no esperar varios minutos y fallar al final.

        Raises:
            ValueError: si los datos no son válidos, si el scraper termina con errores o
                si no guarda ningún precio.
        """
        if Decimal(str(diferencia)) < 0:
            raise ValueError("La diferencia de precios no puede ser negativa.")
        self._servicio_precio.precios_en_pesos(tipo)
        momento = self._ahora()
        codigo = self._lanzador(fecha_extraccion=momento.isoformat())
        if codigo != 0:
            raise ValueError(f"El scraper terminó con errores (código {codigo}).")
        return self.generar_alertas(diferencia, tipo, carpeta, momento)

    @auditar("Generar reporte Excel", detalle=_detalle_reporte)
    def generar_reporte_excel(self, tipo: TipoCotizacion, carpeta: str = CARPETA_EXPORTACION,
                              fecha_extraccion: Optional[datetime.datetime] = None
                              ) -> ResultadoReporte:
        """Genera el reporte en Excel con el precio interno y el precio web de cada libro.

        Hay una fila por libro del catálogo, con las columnas Libro, Precio interno (en
        pesos), Precio web, Diferencia (precio web menos precio interno) y Fecha de
        extracción. Si no hay una publicación que coincida, o el libro no tiene precio
        interno, las celdas que no se pueden calcular quedan vacías.

        Args:
            tipo: tipo de dólar para convertir los precios internos en dólares.
            carpeta: carpeta donde se guarda el archivo.
            fecha_extraccion: corrida del scraper a usar. Por defecto, la más reciente.

        Raises:
            ValueError: si no hay precios web guardados, si falta la cotización del dólar
                o si no está instalada la librería openpyxl.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.comments import Comment
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as error:
            raise ValueError("Falta instalar la librería openpyxl (pip install openpyxl).") \
                from error
        precios_web = self._leer_corrida(fecha_extraccion)
        momento = precios_web[0].fecha_extraccion
        filas = self._comparar(tipo, precios_web)

        libro_excel = Workbook()
        hoja = libro_excel.active
        hoja.title = "Reporte de precios"
        hoja.append(self.COLUMNAS_REPORTE)
        for fila in filas:
            hoja.append([fila.libro.titulo, fila.precio_interno, fila.precio_web,
                         fila.diferencia, fila.fecha_extraccion])

        formato_pesos = '"$" #,##0.00'
        formato_diferencia = '"$" #,##0.00;[Red]-"$" #,##0.00'
        formatos = {2: formato_pesos, 3: formato_pesos, 4: formato_diferencia,
                    5: "dd/mm/yyyy hh:mm"}
        borde = Border(bottom=Side(style="thin", color="808080"))
        for celda in hoja[1]:
            celda.font = Font(name="Arial", size=10, bold=True)
            celda.fill = PatternFill("solid", fgColor="D9E1F2")
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.border = borde
        for fila_excel in hoja.iter_rows(min_row=2):
            for celda in fila_excel:
                celda.font = Font(name="Arial", size=10)
                if celda.column in formatos:
                    celda.number_format = formatos[celda.column]
                    celda.alignment = Alignment(horizontal="right")
        # Aclaraciones donde las ve quien lee el reporte (comentarios de los encabezados).
        aclaraciones = {
            "B1": f"En pesos. Si el libro tiene precio en USD se convirtió con la última "
                  f"venta del dólar {tipo.nombre}.",
            "C1": "Precio de la publicación de Cúspide con el mismo título del libro. "
                  "Vacío si no hay coincidencia.",
            "D1": "Precio web menos precio interno. Positivo: la web es más cara."}
        for referencia, texto in aclaraciones.items():
            comentario = Comment(texto, "Book Manager")
            comentario.width, comentario.height = 280, 90
            hoja[referencia].comment = comentario
        anchos = {"A": min(max((len(f.libro.titulo) for f in filas), default=10) + 2, 60),
                  "B": 16, "C": 16, "D": 16, "E": 20}
        for columna, ancho in anchos.items():
            hoja.column_dimensions[columna].width = max(ancho, 12)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        # Al imprimir o exportar a PDF: hoja horizontal y las cinco columnas en el ancho.
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True

        os.makedirs(carpeta, exist_ok=True)
        ruta = os.path.join(carpeta, f"reporte_precios_{momento:%Y%m%d_%H%M%S}.xlsx")
        libro_excel.save(ruta)
        sin_coincidencia = sum(1 for f in filas if f.precio_web is None)
        return ResultadoReporte(ruta, len(filas), sin_coincidencia, momento)
